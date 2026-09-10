"""DOCX 파일 로딩·이미지 추출 전용 믹스인"""

from typing import List, Dict, Tuple
from langchain.schema import Document
from pathlib import Path
import logging
from datetime import datetime

try:
    from docx import Document as DocxDocument

    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False

logger = logging.getLogger(__name__)


class DocxLoadingMixin:
    """DOCX 파일 로딩·이미지 추출 전용 믹스인 — EnhancedDocumentLoader 믹스인"""

    def _load_docx_file(self, file_path: str, progress_callback=None) -> List[Document]:
        """DOCX 파일 로딩 및 이미지 추출"""
        if not DOCX_AVAILABLE:
            raise ImportError(
                "python-docx 라이브러리가 설치되지 않았습니다. 'pip install python-docx' 명령으로 설치해주세요."
            )

        try:
            document = self._load_docx_document(file_path, progress_callback)
            if progress_callback:
                progress_callback(1.0, "DOCX 처리 완료")
            return [document]
        except Exception as e:
            logger.error(f"DOCX 파일 로딩 실패: {str(e)}")
            # 폴백으로 텍스트 파일로 처리 시도 (일부 내용이라도 추출)
            try:
                logger.info("기본 텍스트 추출 방법으로 재시도...")
                return self._load_text_file(file_path, progress_callback)
            except Exception as fallback_error:
                logger.debug(f"텍스트 폴백도 실패 ({type(fallback_error).__name__}): {fallback_error}")
                raise ValueError(f"DOCX 파일 처리 실패: {str(e)}")

    def _load_docx_document(self, file_path: str, progress_callback) -> Document:
        """DOCX 문서 파싱 → 이미지/텍스트 추출 → Document 생성·MD 저장"""
        file_name = Path(file_path).name
        file_stem = Path(file_path).stem

        if progress_callback:
            progress_callback(0.1, "DOCX 파일 열기 중...")

        doc = DocxDocument(file_path)

        if progress_callback:
            progress_callback(0.3, "텍스트 및 이미지 추출 중...")

        image_count, extracted_images = self._extract_docx_images_safe(doc, file_stem, progress_callback)

        content, paragraph_count, table_count = self._extract_docx_content(
            doc, extracted_images, image_count, progress_callback
        )

        logger.info(f"   📖 DOCX 로딩 완료: {paragraph_count}개 단락, {table_count}개 표, {image_count}개 이미지")

        document = self._build_docx_document(
            file_path, file_name, content, paragraph_count, table_count, image_count, extracted_images
        )
        self._save_docx_markdown(document, file_path, file_name, content, paragraph_count, table_count, image_count)

        return document

    def _extract_docx_content(
        self, doc, extracted_images: List[Dict], image_count: int, progress_callback
    ) -> Tuple[str, int, int]:
        """단락·표·이미지 참조를 결합해 최종 마크다운 콘텐츠 생성"""
        if progress_callback:
            progress_callback(0.5, "텍스트 추출 중...")

        full_text, paragraph_count, table_count = self._extract_docx_text(doc)

        if progress_callback:
            progress_callback(0.7, f"텍스트 정리 중... ({paragraph_count}개 단락)")

        full_text.extend(self._build_docx_image_references(extracted_images))

        if progress_callback:
            progress_callback(0.9, f"처리 완료 ({paragraph_count}개 단락, {table_count}개 표, {image_count}개 이미지)")

        # 전체 내용 결합
        content = "\n\n".join(full_text)

        if not content.strip():
            raise ValueError("DOCX 파일에서 텍스트를 추출할 수 없습니다.")

        return content, paragraph_count, table_count

    def _extract_docx_images_safe(self, doc, file_stem: str, progress_callback) -> Tuple[int, List[Dict]]:
        """이미지 추출(실패 시 경고 후 빈 결과)"""
        images_dir = Path("static/images/docx")
        images_dir.mkdir(parents=True, exist_ok=True)

        try:
            image_count, extracted_images = self._extract_docx_images(doc, file_stem, images_dir)
            if progress_callback:
                progress_callback(0.4, f"이미지 추출 완료 ({image_count}개)")
            return image_count, extracted_images
        except Exception as e:
            logger.warning(f"DOCX 이미지 추출 중 오류: {str(e)}")
            return 0, []

    def _extract_docx_text(self, doc) -> Tuple[List[str], int, int]:
        """단락과 표에서 텍스트 블록 추출"""
        full_text = []
        paragraph_count = 0

        # 단락별로 텍스트 추출
        for paragraph in doc.paragraphs:
            text = paragraph.text.strip()
            if text:  # 빈 단락 제외
                full_text.append(text)
                paragraph_count += 1

        # 표(table) 내용도 추출
        table_count = 0
        for table in doc.tables:
            table_text = self._render_docx_table(table)
            if table_text:
                full_text.append(table_text)
                table_count += 1

        return full_text, paragraph_count, table_count

    @staticmethod
    def _render_docx_table(table) -> str:
        """단일 표를 파이프(|) 구분 텍스트로 렌더링"""
        table_text = []
        for row in table.rows:
            row_text = []
            for cell in row.cells:
                cell_text = cell.text.strip()
                if cell_text:
                    row_text.append(cell_text)
            if row_text:
                table_text.append(" | ".join(row_text))

        return "\n".join(table_text)

    @staticmethod
    def _build_docx_image_references(extracted_images: List[Dict]) -> List[str]:
        """추출된 이미지의 마크다운 참조 블록 생성"""
        if not extracted_images:
            return []

        refs = ["\n## 추출된 이미지"]
        for i, image_info in enumerate(extracted_images, 1):
            refs.append(f"![이미지 {i}]({image_info['relative_path']})")
        return refs

    @staticmethod
    def _build_docx_document(
        file_path: str,
        file_name: str,
        content: str,
        paragraph_count: int,
        table_count: int,
        image_count: int,
        extracted_images: List[Dict],
    ) -> Document:
        """DOCX 로딩 결과 Document 객체 생성"""
        return Document(
            page_content=content,
            metadata={
                "source": str(file_path),
                "file_name": file_name,
                "file_type": ".docx",
                "paragraph_count": paragraph_count,
                "table_count": table_count,
                "image_count": image_count,
                "images": extracted_images,
                "original_size": len(content),
                "processing_method": "docx_with_images",
            },
        )

    def _save_docx_markdown(
        self,
        document: Document,
        file_path: str,
        file_name: str,
        content: str,
        paragraph_count: int,
        table_count: int,
        image_count: int,
    ) -> None:
        """MD 파일 저장 (전처리 확인용)"""
        try:
            docx_name = Path(file_path).stem
            md_file_path = Path("converted_docs") / f"{docx_name}.md"
            md_file_path.parent.mkdir(parents=True, exist_ok=True)

            with open(md_file_path, "w", encoding="utf-8") as f:
                f.write(self._docx_md_header(docx_name, file_name, paragraph_count, table_count, image_count))
                f.write(content)

            # 메타데이터에 MD 파일 경로 추가
            document.metadata["md_file_path"] = str(md_file_path)
            document.metadata["md_saved"] = True

            logger.info(f"   📝 MD 파일 저장: {md_file_path}")

        except Exception as e:
            logger.warning(f"MD 파일 저장 실패: {str(e)}")
            document.metadata["md_saved"] = False

    @staticmethod
    def _docx_md_header(
        docx_name: str, file_name: str, paragraph_count: int, table_count: int, image_count: int
    ) -> str:
        """MD 변환 헤더 메타 블록 생성"""
        return (
            f"# {docx_name}\n\n"
            f"**원본 파일**: {file_name}\n"
            f"**변환 시간**: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n"
            f"**단락 수**: {paragraph_count}개\n"
            f"**표 수**: {table_count}개\n"
            f"**추출된 이미지**: {image_count}개\n\n---\n\n"
        )

    def _extract_docx_images(self, doc, file_stem: str, images_dir: Path) -> Tuple[int, List[Dict]]:
        """
        DOCX 문서에서 이미지를 추출하고 저장

        Args:
            doc: python-docx Document 객체
            file_stem: 파일명 (확장자 제외)
            images_dir: 이미지 저장 디렉토리

        Returns:
            Tuple[이미지 개수, 이미지 정보 리스트]
        """

        extracted_images = []
        image_count = 0

        try:
            # DOCX 파일의 모든 관련 parts 확인
            for rel in doc.part.rels.values():
                if "image" not in rel.target_ref:
                    continue
                try:
                    image_info = self._extract_single_docx_image(rel, file_stem, images_dir, image_count)
                    extracted_images.append(image_info)
                    image_count += 1
                except Exception as e:
                    logger.warning(f"개별 이미지 추출 실패: {str(e)}")
                    continue

            # 인라인 이미지도 확인 (inline shapes)
            try:
                for shape in doc.inline_shapes:
                    if hasattr(shape, "_inline") and hasattr(shape._inline, "graphic"):
                        # 인라인 이미지 처리 (추가 구현 가능)
                        pass
            except Exception as e:
                logger.debug(f"인라인 이미지 처리 중 오류: {str(e)}")

        except Exception as e:
            logger.error(f"DOCX 이미지 추출 실패: {str(e)}")
            return 0, []

        return image_count, extracted_images

    def _extract_single_docx_image(self, rel, file_stem: str, images_dir: Path, index: int) -> Dict:
        """개별 이미지 관계를 파일로 저장하고 메타데이터 반환"""
        from datetime import datetime

        image_part = rel.target_part
        image_data = image_part.blob
        content_type = getattr(image_part, "content_type", "")

        ext = self._guess_docx_image_extension(content_type, image_data)
        image_filename = f"{file_stem}_img{index + 1}{ext}"
        image_path = images_dir / image_filename

        # 이미지 저장
        with open(image_path, "wb") as f:
            f.write(image_data)

        logger.info(f"   🖼️  DOCX 이미지 추출: {image_filename} ({len(image_data)} bytes)")

        return {
            "filename": image_filename,
            "path": str(image_path),
            "relative_path": f"static/images/docx/{image_filename}",
            "size": len(image_data),
            "format": ext[1:].upper(),
            "content_type": content_type,
            "extracted_at": datetime.now().isoformat(),
        }

    @staticmethod
    def _guess_docx_image_extension(content_type: str, image_data: bytes) -> str:
        """content_type 우선, 없으면 이미지 시그니처로 확장자 추정"""
        if "jpeg" in content_type or "jpg" in content_type:
            return ".jpg"
        if "png" in content_type:
            return ".png"
        if "gif" in content_type:
            return ".gif"
        if "bmp" in content_type:
            return ".bmp"

        # 이미지 시그니처로 확장자 추정
        if image_data.startswith(b"\xff\xd8"):
            return ".jpg"
        if image_data.startswith(b"\x89PNG"):
            return ".png"
        if image_data.startswith(b"GIF8"):
            return ".gif"
        if image_data.startswith(b"BM"):
            return ".bmp"
        return ".png"  # 기본값
